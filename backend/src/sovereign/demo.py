"""Two synthetic companies for the demo and the security tests.

Musterwerk GmbH has teams with different access: engineers can't read HR's salary bands,
and finance documents are limited to finance. RheinTech GmbH is a second tenant whose
documents deliberately overlap with Musterwerk's (both have an error code E104 and a remote
work rule, with different answers), so any cross-tenant leak is visible in the answer.
Felix is a consultant who belongs to both organizations.
"""

import io
import uuid
from dataclasses import dataclass, field
from pathlib import Path

import docx
from sqlalchemy import select, text
from sqlalchemy.orm import Session, sessionmaker

from sovereign.ai.embeddings import Embedder
from sovereign.auth import hash_password
from sovereign.config import Settings
from sovereign.db import system_session
from sovereign.ingestion import pipeline
from sovereign.ingestion.extract import detect_type
from sovereign.models import (
    Classification,
    Document,
    DocumentVersion,
    KnowledgeSpace,
    Membership,
    Permission,
    PrincipalType,
    Role,
    SpaceGrant,
    Team,
    TeamMember,
    User,
)
from sovereign.organizations import create_organization
from sovereign.storage import Storage

PASSWORD = "demo-password-2026"
DOCS = Path(__file__).resolve().parents[3] / "demo-data" / "documents"

USERS = {
    "anna": ("anna.schmidt@musterwerk.example", "Anna Schmidt"),
    "ben": ("ben.weber@musterwerk.example", "Ben Weber"),
    "clara": ("clara.fischer@musterwerk.example", "Clara Fischer"),
    "david": ("david.koch@musterwerk.example", "David Koch"),
    "eva": ("eva.braun@rheintech.example", "Eva Braun"),
    "felix": ("felix.wagner@consulting.example", "Felix Wagner"),
}


@dataclass
class DemoIds:
    users: dict[str, uuid.UUID] = field(default_factory=dict)
    orgs: dict[str, uuid.UUID] = field(default_factory=dict)
    memberships: dict[tuple[str, str], uuid.UUID] = field(default_factory=dict)
    teams: dict[str, uuid.UUID] = field(default_factory=dict)
    spaces: dict[str, uuid.UUID] = field(default_factory=dict)
    documents: dict[str, uuid.UUID] = field(default_factory=dict)


def remote_work_policy_docx() -> bytes:
    """Built at seed time so the demo exercises the DOCX parser."""
    document = docx.Document()
    document.add_heading("Remote Work Policy 2026", level=1)
    document.add_paragraph(
        "Synthetic demonstration document for the fictional Musterwerk GmbH. "
        "Valid from 1 January 2026."
    )
    document.add_heading("1 Eligibility", level=2)
    document.add_paragraph(
        "All employees who have completed their probation period may work remotely. "
        "Roles that require presence at machines (workshop, assembly) are excluded."
    )
    document.add_heading("2 Remote Days", level=2)
    document.add_paragraph(
        "Employees may work remotely up to three days per week. Tuesday is a team day "
        "and everyone works on site. Remote work must be done from Germany; working from "
        "another EU country needs HR approval and is limited to 20 days per year."
    )
    document.add_heading("3 Equipment Allowance", level=2)
    document.add_paragraph(
        "Musterwerk pays a one-time home office allowance of 300 euros for a monitor and "
        "chair. Receipts are submitted in the travel tool."
    )
    document.add_heading("4 Availability", level=2)
    document.add_paragraph(
        "During core hours (10:00 to 15:00) remote employees must be reachable by phone and chat."
    )
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def reset(factory: sessionmaker[Session]) -> None:
    """Delete all data. Development only: runs as the login role, which the app never does."""
    tables = ", ".join(t.name for t in reversed(Document.metadata.sorted_tables))
    with factory.kw["bind"].begin() as connection:
        connection.execute(text(f"TRUNCATE {tables} CASCADE"))


def seed(
    factory: sessionmaker[Session],
    settings: Settings,
    storage: Storage,
    embedder: Embedder,
    *,
    docs_dir: Path = DOCS,
    run_ingestion: bool = True,
) -> DemoIds:
    ids = DemoIds()
    with system_session(factory) as session:
        if session.scalar(select(User).where(User.email == USERS["anna"][0])):
            raise RuntimeError("Demo data already exists. Run with --reset to recreate it.")
        password_hash = hash_password(PASSWORD)
        users = {}
        for key, (email, name) in USERS.items():
            users[key] = User(email=email, display_name=name, password_hash=password_hash)
            session.add(users[key])
        session.flush()
        ids.users = {key: user.id for key, user in users.items()}

        def member(org, key: str, role: Role) -> Membership:
            m = Membership(organization_id=org.id, user_id=users[key].id, role=role)
            session.add(m)
            session.flush()
            ids.memberships[(org.slug, key)] = m.id
            return m

        def space(org, name: str, description: str, classification=Classification.INTERNAL):
            s = KnowledgeSpace(
                organization_id=org.id,
                name=name,
                description=description,
                classification=classification,
            )
            session.add(s)
            session.flush()
            ids.spaces[f"{org.slug}/{name}"] = s.id
            return s

        def grant(s, principal_type, principal_id, permission=Permission.READ) -> None:
            session.add(
                SpaceGrant(
                    organization_id=s.organization_id,
                    space_id=s.id,
                    principal_type=principal_type,
                    principal_id=principal_id,
                    permission=permission,
                )
            )

        def team(org, name: str, members: list[Membership]) -> Team:
            t = Team(organization_id=org.id, name=name)
            session.add(t)
            session.flush()
            for m in members:
                session.add(TeamMember(organization_id=org.id, team_id=t.id, membership_id=m.id))
            ids.teams[f"{org.slug}/{name}"] = t.id
            return t

        def add_document(s, title: str, filename: str, data: bytes, uploader: Membership):
            d = Document(
                organization_id=s.organization_id,
                space_id=s.id,
                title=title,
                created_by=uploader.id,
            )
            session.add(d)
            session.flush()
            key, checksum = storage.put(s.organization_id, data)
            v = DocumentVersion(
                organization_id=s.organization_id,
                document_id=d.id,
                version_no=1,
                filename=filename,
                mime_type=detect_type(data, filename),
                size_bytes=len(data),
                checksum=checksum,
                storage_key=key,
                uploaded_by=uploader.id,
            )
            session.add(v)
            session.flush()
            pipeline.enqueue(session, v)
            ids.documents[title] = d.id

        def read(*parts: str) -> bytes:
            return docs_dir.joinpath(*parts).read_bytes()

        # Musterwerk GmbH
        mw, anna = create_organization(
            session, settings, name="Musterwerk GmbH", slug="musterwerk", owner=users["anna"]
        )
        ids.orgs["musterwerk"] = mw.id
        ids.memberships[("musterwerk", "anna")] = anna.id
        ben = member(mw, "ben", Role.EMPLOYEE)
        clara = member(mw, "clara", Role.EDITOR)
        david = member(mw, "david", Role.EMPLOYEE)
        member(mw, "felix", Role.EMPLOYEE)
        engineering = team(mw, "Engineering", [ben])
        hr = team(mw, "HR", [clara])
        finance = team(mw, "Finance", [david])

        handbook = session.scalar(
            select(KnowledgeSpace).where(
                KnowledgeSpace.organization_id == mw.id, KnowledgeSpace.name == "Company Handbook"
            )
        )
        ids.spaces["musterwerk/Company Handbook"] = handbook.id
        grant(handbook, PrincipalType.TEAM, hr.id, Permission.EDIT)
        eng_space = space(mw, "Engineering", "Manuals and engineering processes.")
        grant(eng_space, PrincipalType.TEAM, engineering.id)
        grant(eng_space, PrincipalType.MEMBER, anna.id, Permission.EDIT)
        hr_space = space(
            mw, "HR Confidential", "Salary data. HR only.", Classification.CONFIDENTIAL
        )
        grant(hr_space, PrincipalType.TEAM, hr.id, Permission.EDIT)
        fin_space = space(mw, "Finance", "Procurement and supplier contracts.")
        grant(fin_space, PrincipalType.TEAM, finance.id)
        grant(fin_space, PrincipalType.MEMBER, anna.id, Permission.EDIT)

        add_document(
            handbook,
            "Employee Handbook 2026",
            "Employee_Handbook_2026.md",
            read("musterwerk", "handbook", "Employee_Handbook_2026.md"),
            clara,
        )
        add_document(
            handbook,
            "Remote Work Policy 2026",
            "Remote_Work_Policy_2026.docx",
            remote_work_policy_docx(),
            clara,
        )
        add_document(
            handbook,
            "Travel Reimbursement Rules",
            "Travel_Reimbursement_Rules.md",
            read("musterwerk", "handbook", "Travel_Reimbursement_Rules.md"),
            clara,
        )
        add_document(
            handbook,
            "Customer Escalation Procedure",
            "Customer_Escalation_Procedure.md",
            read("musterwerk", "handbook", "Customer_Escalation_Procedure.md"),
            anna,
        )
        add_document(
            eng_space,
            "MW-400 Troubleshooting Manual",
            "MW-400_Troubleshooting_Manual.md",
            read("musterwerk", "engineering", "MW-400_Troubleshooting_Manual.md"),
            anna,
        )
        add_document(
            eng_space,
            "Engineering Onboarding Guide",
            "Engineering_Onboarding.md",
            read("musterwerk", "engineering", "Engineering_Onboarding.md"),
            anna,
        )
        add_document(
            hr_space,
            "Salary Bands 2026",
            "Salary_Bands_2026.md",
            read("musterwerk", "hr", "Salary_Bands_2026.md"),
            clara,
        )
        for title, filename in [
            ("Procurement Policy 2025", "Procurement_Policy_2025.md"),
            ("Procurement Policy 2026", "Procurement_Policy_2026.md"),
            ("Supplier Agreement Rheinstahl", "Supplier_Agreement_Rheinstahl.md"),
        ]:
            add_document(fin_space, title, filename, read("musterwerk", "finance", filename), anna)

        # RheinTech GmbH
        rt, eva = create_organization(
            session, settings, name="RheinTech GmbH", slug="rheintech", owner=users["eva"]
        )
        ids.orgs["rheintech"] = rt.id
        ids.memberships[("rheintech", "eva")] = eva.id
        member(rt, "felix", Role.EMPLOYEE)
        rt_handbook = session.scalar(
            select(KnowledgeSpace).where(
                KnowledgeSpace.organization_id == rt.id, KnowledgeSpace.name == "Company Handbook"
            )
        )
        ids.spaces["rheintech/Company Handbook"] = rt_handbook.id
        rt_service = space(rt, "Service", "Service manuals.")
        grant(rt_service, PrincipalType.ORGANIZATION, rt.id)
        grant(rt_service, PrincipalType.MEMBER, eva.id, Permission.EDIT)
        add_document(
            rt_handbook,
            "RheinTech Handbook",
            "RheinTech_Handbook.md",
            read("rheintech", "RheinTech_Handbook.md"),
            eva,
        )
        add_document(
            rt_service,
            "RT-Pump Service Manual",
            "RT-Pump_Service_Manual.md",
            read("rheintech", "RT-Pump_Service_Manual.md"),
            eva,
        )
        session.commit()

    if run_ingestion:
        pipeline.run_pending(factory, storage, embedder)
    return ids
