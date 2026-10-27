import io

import docx
import pytest

from sovereign.ingestion.chunking import chunk_sections
from sovereign.ingestion.extract import (
    DOCX,
    MARKDOWN,
    PDF,
    TEXT,
    ExtractionError,
    Section,
    UnsupportedFile,
    detect_type,
    extract,
)


def make_pdf(pages: list[str]) -> bytes:
    """A minimal valid PDF with one line of Helvetica text per page."""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", ""]
    kids = []
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        content_id = len(objects) + 1
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        page_id = len(objects) + 1
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R"
            " /Resources << /Font << /F1 << /Type /Font /Subtype /Type1"
            " /BaseFont /Helvetica >> >> >> >>"
        )
        kids.append(f"{page_id} 0 R")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>"

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n{body}\nendobj\n".encode())
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    )
    return out.getvalue()


def make_docx() -> bytes:
    document = docx.Document()
    document.add_heading("Remote Work Policy", level=1)
    document.add_paragraph("This policy applies to all employees.")
    document.add_heading("Eligibility", level=2)
    document.add_paragraph("Employees may work remotely up to three days per week.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Role"
    table.cell(0, 1).text = "Days"
    table.cell(1, 0).text = "Engineer"
    table.cell(1, 1).text = "3"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def test_detects_types_from_content_not_extension():
    assert detect_type(make_pdf(["hi"]), "policy.txt") == PDF
    assert detect_type(make_docx(), "policy.pdf") == DOCX
    assert detect_type(b"# Title", "notes.md") == MARKDOWN
    assert detect_type(b"plain", "notes.txt") == TEXT
    with pytest.raises(UnsupportedFile):
        detect_type(b"\x89PNG\r\n\x1a\n\x00\x00", "image.md")
    with pytest.raises(UnsupportedFile):
        detect_type(b"PK\x03\x04garbage", "archive.docx")


def test_pdf_sections_carry_page_numbers():
    sections = extract(make_pdf(["First page text", "Second page text"]), PDF)
    assert [s.location for s in sections] == [{"page": 1}, {"page": 2}]
    assert "Second page" in sections[1].text


def test_pdf_without_text_is_rejected():
    with pytest.raises(ExtractionError, match="OCR"):
        extract(make_pdf([""]), PDF)


def test_docx_sections_follow_headings_and_include_tables():
    sections = extract(make_docx(), DOCX)
    eligibility = next(s for s in sections if "three days" in s.text)
    assert eligibility.location == {
        "section": "Eligibility",
        "path": ["Remote Work Policy", "Eligibility"],
    }
    assert any("Engineer | 3" in s.text for s in sections)


def test_markdown_heading_path_and_code_blocks():
    text = (
        "# Handbook\nIntro.\n## 4 Leave\n### 4.2 Annual Leave\n30 days.\n"
        "```\n# not a heading\n```\n## 5 Travel\nTrain first."
    )
    sections = extract(text.encode(), MARKDOWN)
    leave = next(s for s in sections if "30 days" in s.text)
    assert leave.location["path"] == ["Handbook", "4 Leave", "4.2 Annual Leave"]
    assert "# not a heading" in leave.text
    travel = next(s for s in sections if "Train" in s.text)
    assert travel.location["path"] == ["Handbook", "5 Travel"]


def test_chunks_stay_within_sections_and_carry_headers():
    sections = [
        Section("A" * 50, {"section": "One", "path": ["One"]}),
        Section("\n\n".join(["word " * 100] * 6), {"section": "Two", "path": ["Two"]}),
    ]
    chunks = chunk_sections("Handbook", sections, max_chars=1200)
    assert chunks[0].context_header == "Handbook > One"
    assert all(c.context_header == "Handbook > Two" for c in chunks[1:])
    assert len(chunks) > 2
    assert all(len(c.content) <= 1200 + 2 for c in chunks)
    assert [c.index for c in chunks] == list(range(len(chunks)))


def test_oversized_paragraph_is_split_at_sentences():
    paragraph = " ".join(f"Sentence number {i} is here." for i in range(200))
    chunks = chunk_sections("Doc", [Section(paragraph)], max_chars=500)
    assert len(chunks) > 1
    assert all(len(c.content) <= 500 for c in chunks)
    assert chunks[0].content.endswith(".")
